import docx
from typing import Dict, Any, List


class DocxDocumentParser:
    """DOCX document parsing utility using python-docx."""

    def extract_paragraphs_and_tables(self, file_path: str) -> Dict[str, Any]:
        """Extract paragraph text and table contents from a DOCX file."""
        doc = docx.Document(file_path)
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        tables_data = []

        for table in doc.tables:
            table_rows = []
            for row in table.rows:
                table_rows.append([cell.text.strip() for cell in row.cells])
            tables_data.append(table_rows)

        return {
            "paragraphs": paragraphs,
            "tables": tables_data
        }

# TODO: Preserve heading structure and inline images from DOCX
